from __future__ import annotations

import math
import re
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor


DOCS = Path(__file__).resolve().parent
ROOT = DOCS.parent
ASSETS = DOCS / "report-assets"
MARKDOWN = DOCS / "Bao_Cao_Tieu_Luan_TMDT.md"
OUTPUT = DOCS / "Bao_Cao_Tieu_Luan_TMDT.docx"

PAPER = "#F8F7F2"
SURFACE = "#FFFEF9"
JADE = "#0B513C"
JADE_DARK = "#073F30"
JADE_LIGHT = "#DCE8DC"
INK = "#10251E"
MUTED = "#52655D"
LINE = "#AFC2B6"
ORANGE = "#C64A0A"


def hex_rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "timesbd.ttf" if bold else "times.ttf"
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / name), size=size)


def fit_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), trial, font=font)[2] <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered(draw: ImageDraw.ImageDraw, text: str, box: tuple[int, int, int, int], font, fill=INK, spacing=8):
    x1, y1, x2, y2 = box
    lines = fit_text(draw, text, font, x2 - x1 - 36)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + spacing * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=font)[2]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=font, fill=hex_rgb(fill))
        y += height + spacing


def rounded_box(draw, box, text, *, fill=SURFACE, outline=LINE, font=None, radius=26, text_fill=INK, width=4):
    draw.rounded_rectangle(box, radius=radius, fill=hex_rgb(fill), outline=hex_rgb(outline), width=width)
    draw_centered(draw, text, box, font or load_font(34, True), text_fill)


def arrow(draw, start, end, *, fill=JADE, width=7):
    draw.line([start, end], fill=hex_rgb(fill), width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 22
    wing = 0.55
    p1 = (end[0] - length * math.cos(angle - wing), end[1] - length * math.sin(angle - wing))
    p2 = (end[0] - length * math.cos(angle + wing), end[1] - length * math.sin(angle + wing))
    draw.polygon([end, p1, p2], fill=hex_rgb(fill))


def diagram_canvas(title: str, subtitle: str, size=(2400, 1500)):
    img = Image.new("RGB", size, hex_rgb(PAPER))
    draw = ImageDraw.Draw(img)
    draw.text((90, 70), title, font=load_font(58, True), fill=hex_rgb(JADE_DARK))
    draw.text((92, 145), subtitle, font=load_font(30), fill=hex_rgb(MUTED))
    draw.line((90, 205, size[0] - 90, 205), fill=hex_rgb(JADE_LIGHT), width=5)
    return img, draw


def save_diagram(img: Image.Image, name: str):
    img.save(ASSETS / name, format="PNG", optimize=True)


def make_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)

    img, d = diagram_canvas("Use case tổng quát", "Phạm vi được xác nhận từ storefront và trang quản trị")
    rounded_box(d, (90, 360, 430, 610), "KHÁCH HÀNG", fill=JADE_DARK, outline=JADE_DARK, text_fill=SURFACE, font=load_font(38, True))
    rounded_box(d, (1970, 360, 2310, 610), "ADMIN /\nOPERATOR", fill=JADE_DARK, outline=JADE_DARK, text_fill=SURFACE, font=load_font(38, True))
    customer = ["Xem và tìm sản phẩm", "Tùy chỉnh đồ uống", "Quản lý giỏ hàng", "Nhập giao hàng", "Đặt đơn mô phỏng"]
    admin = ["Đăng nhập nội bộ", "Xem dashboard và đơn", "Lọc, tìm, phân trang", "Chuyển trạng thái", "Xác nhận COD"]
    ys = [280, 490, 700, 910, 1120]
    for text, y in zip(customer, ys):
        rounded_box(d, (630, y, 1190, y + 145), text, font=load_font(31, True))
        arrow(d, (430, 485), (630, y + 72), fill=JADE)
    for text, y in zip(admin, ys):
        rounded_box(d, (1240, y, 1800, y + 145), text, font=load_font(31, True))
        arrow(d, (1970, 485), (1800, y + 72), fill=JADE)
    d.text((650, 1340), "Ngoài phạm vi hiện tại: tài khoản khách và tra cứu trạng thái đơn công khai", font=load_font(29), fill=hex_rgb(ORANGE))
    save_diagram(img, "use-case.png")

    img, d = diagram_canvas("Luồng nghiệp vụ đặt hàng", "Các nhánh lỗi giữ nguyên dữ liệu để người dùng sửa hoặc nhập tay")
    steps = ["Chọn món", "Tùy chỉnh", "Giỏ hàng", "Thông tin giao", "Server kiểm tra\nvà tính lại giá", "Tạo pending", "Xác nhận mô phỏng", "Nhận mã TSN"]
    coords = []
    for i, text in enumerate(steps):
        row, col = divmod(i, 4)
        x = 100 + col * 565
        y = 330 + row * 480
        coords.append((x, y))
        rounded_box(d, (x, y, x + 430, y + 180), text, fill=SURFACE if i not in (4, 5) else JADE_LIGHT, outline=JADE, font=load_font(31, True))
    for i in range(3):
        arrow(d, (coords[i][0] + 430, coords[i][1] + 90), (coords[i + 1][0], coords[i + 1][1] + 90))
    arrow(d, (coords[3][0] + 215, coords[3][1] + 180), (coords[4][0] + 215, coords[4][1]))
    for i in range(4, 7):
        arrow(d, (coords[i][0] + 430, coords[i][1] + 90), (coords[i + 1][0], coords[i + 1][1] + 90))
    d.text((120, 1290), "Geoapify lỗi / không có kết quả → nhập địa chỉ thủ công", font=load_font(30), fill=hex_rgb(MUTED))
    d.text((1200, 1290), "Validation lỗi → hiển thị trường cần sửa, không mất form", font=load_font(30), fill=hex_rgb(MUTED))
    save_diagram(img, "order-flow.png")

    img, d = diagram_canvas("Kiến trúc tổng thể", "Vinext/Vite application trên Cloudflare Worker, dữ liệu vận hành trong D1")
    layers = [
        ("TRÌNH DUYỆT", ["Storefront React", "Admin React"]),
        ("HTTP ROUTE", ["API menu / địa chỉ / đơn", "API auth / dashboard / orders"]),
        ("NGHIỆP VỤ", ["Order service", "Admin service", "Auth service", "Address service"]),
        ("ADAPTER VÀ DỮ LIỆU", ["D1 repository", "PaymentProvider", "Catalogue nội bộ", "Geoapify"]),
    ]
    y = 280
    for li, (label, boxes) in enumerate(layers):
        d.text((100, y + 55), label, font=load_font(28, True), fill=hex_rgb(JADE_DARK))
        start_x = 430
        available = 1850
        gap = 35
        bw = int((available - gap * (len(boxes) - 1)) / len(boxes))
        for j, text in enumerate(boxes):
            x = start_x + j * (bw + gap)
            rounded_box(d, (x, y, x + bw, y + 150), text, fill=JADE_LIGHT if li == 2 else SURFACE, outline=JADE, font=load_font(29, True))
        if li < len(layers) - 1:
            arrow(d, (1330, y + 150), (1330, y + 245), fill=ORANGE)
        y += 285
    save_diagram(img, "architecture.png")

    img, d = diagram_canvas("Module và seam", "Interface ổn định giúp thay adapter mà không đổi luồng cốt lõi")
    centers = {
        "UI khách": (320, 420), "Checkout API": (900, 420), "Order service": (1510, 420), "Repository": (2080, 420),
        "UI admin": (320, 920), "Admin API": (900, 920), "Admin service": (1510, 920), "D1": (2080, 920),
    }
    for label, (cx, cy) in centers.items():
        fill = JADE_LIGHT if "service" in label.lower() else SURFACE
        rounded_box(d, (cx - 210, cy - 90, cx + 210, cy + 90), label, fill=fill, outline=JADE, font=load_font(31, True))
    for a, b in [("UI khách", "Checkout API"), ("Checkout API", "Order service"), ("Order service", "Repository"), ("Repository", "D1"), ("UI admin", "Admin API"), ("Admin API", "Admin service"), ("Admin service", "D1")]:
        x1, y1 = centers[a]; x2, y2 = centers[b]
        arrow(d, (x1 + 210, y1), (x2 - 210, y2))
    rounded_box(d, (1295, 1190, 1725, 1360), "PaymentProvider", fill=SURFACE, outline=ORANGE, font=load_font(30, True))
    arrow(d, (1510, 510), (1510, 1190), fill=ORANGE)
    d.text((1770, 1245), "COD / QR mô phỏng\n→ cổng thật trong tương lai", font=load_font(29), fill=hex_rgb(MUTED))
    save_diagram(img, "module-seams.png")

    img, d = diagram_canvas("Sequence quy trình đặt hàng", "Giá được xác định ở server; xác nhận không đồng nghĩa đã thanh toán")
    actors = ["Khách", "Storefront", "Order API", "Order service", "D1 repository", "PaymentProvider"]
    xs = [160, 570, 980, 1390, 1800, 2210]
    for x, actor in zip(xs, actors):
        rounded_box(d, (x - 150, 260, x + 150, 365), actor, fill=JADE_DARK, outline=JADE_DARK, text_fill=SURFACE, font=load_font(25, True), radius=18)
        d.line((x, 365, x, 1390), fill=hex_rgb(LINE), width=4)
    events = [
        (0, 1, 470, "1. Gửi form + items"), (1, 2, 590, "2. POST /api/orders"),
        (2, 3, 710, "3. Validate payload"), (3, 5, 830, "4. prepare(total)"),
        (3, 4, 950, "5. Ghi order pending"), (4, 1, 1070, "6. Trả order + payment"),
        (1, 2, 1190, "7. POST /confirm"), (2, 4, 1310, "8. confirmed; QR vẫn simulation_only"),
    ]
    for a, b, y, label in events:
        direction = 1 if xs[b] > xs[a] else -1
        arrow(d, (xs[a] + direction * 10, y), (xs[b] - direction * 10, y), fill=ORANGE if y in (830, 1310) else JADE, width=5)
        mid = (xs[a] + xs[b]) / 2
        tw = d.textbbox((0, 0), label, font=load_font(24))[2]
        d.rectangle((mid - tw / 2 - 8, y - 35, mid + tw / 2 + 8, y - 5), fill=hex_rgb(PAPER))
        d.text((mid - tw / 2, y - 36), label, font=load_font(24), fill=hex_rgb(INK))
    save_diagram(img, "order-sequence.png")

    img, d = diagram_canvas("Mô hình quan hệ dữ liệu", "Chín bảng D1, tiền lưu bằng số nguyên VND")
    tables = {
        "orders": (930, 260, ["PK id", "order_code UQ", "customer + address", "amounts + status", "version"]),
        "order_items": (380, 610, ["PK id", "FK order_id", "product snapshot", "options + quantity"]),
        "payments": (1480, 610, ["PK id", "FK order_id UQ", "provider + status", "due / paid"]),
        "order_item_toppings": (170, 1030, ["PK id", "FK order_item_id", "topping snapshot"]),
        "order_status_history": (760, 1030, ["PK id", "FK order_id", "FK actor_admin_id", "from / to / reason"]),
        "admin_audit_logs": (1350, 1030, ["PK id", "FK admin_id", "before / after", "request metadata"]),
        "admin_users": (1920, 1030, ["PK id", "login + role", "password hash", "lock state"]),
        "admin_sessions": (1950, 560, ["PK id", "FK admin_id", "token + CSRF hash", "expires / revoked"]),
        "admin_login_attempts": (1950, 260, ["PK key_hash", "attempt count", "window / blocked"]),
    }
    boxes = {}
    for name, (cx, cy, fields) in tables.items():
        w, h = 390, 235
        box = (cx - w // 2, cy, cx + w // 2, cy + h)
        boxes[name] = box
        d.rounded_rectangle(box, 18, fill=hex_rgb(SURFACE), outline=hex_rgb(JADE), width=4)
        d.rectangle((box[0], box[1], box[2], box[1] + 58), fill=hex_rgb(JADE))
        d.text((box[0] + 18, box[1] + 10), name, font=load_font(27, True), fill=hex_rgb(SURFACE))
        yy = box[1] + 72
        for field in fields:
            d.text((box[0] + 18, yy), field, font=load_font(22), fill=hex_rgb(INK)); yy += 34
    def connect(a, b, color=JADE):
        ba, bb = boxes[a], boxes[b]
        sa = ((ba[0] + ba[2]) // 2, ba[3]); eb = ((bb[0] + bb[2]) // 2, bb[1])
        arrow(d, sa, eb, fill=color, width=4)
    connect("orders", "order_items"); connect("orders", "payments"); connect("order_items", "order_item_toppings")
    connect("orders", "order_status_history"); connect("admin_users", "admin_sessions"); connect("admin_users", "admin_audit_logs")
    connect("admin_users", "order_status_history", ORANGE)
    save_diagram(img, "erd.png")

    img, d = diagram_canvas("State machine đơn hàng", "Transition sai hoặc version cũ trả lỗi 409")
    nodes = {
        "pending": (250, 560), "confirmed": (720, 560), "preparing": (1190, 560),
        "delivering": (1660, 560), "completed": (2130, 560),
        "rejected": (560, 1080), "cancelled": (1400, 1080),
    }
    for name, (x, y) in nodes.items():
        final = name in {"completed", "cancelled", "rejected"}
        rounded_box(d, (x - 170, y - 80, x + 170, y + 80), name, fill=JADE_DARK if final else SURFACE, outline=JADE, text_fill=SURFACE if final else INK, font=load_font(31, True))
    for a, b in [("pending", "confirmed"), ("confirmed", "preparing"), ("preparing", "delivering"), ("delivering", "completed")]:
        ax, ay = nodes[a]; bx, by = nodes[b]; arrow(d, (ax + 170, ay), (bx - 170, by))
    arrow(d, (250, 640), (510, 1000), fill=ORANGE)
    for a in ["pending", "confirmed", "preparing", "delivering"]:
        ax, ay = nodes[a]
        arrow(d, (ax, ay + 80), (1400, 1000), fill=ORANGE, width=4)
    d.text((835, 1290), "Trạng thái cuối không có transition tiếp theo", font=load_font(32, True), fill=hex_rgb(MUTED))
    save_diagram(img, "order-state-machine.png")

    screenshots = {
        ROOT / ".impeccable/review/tablet-768.png": ASSETS / "storefront-768.png",
        ROOT / ".impeccable/review/tablet-1024.png": ASSETS / "storefront-1024.png",
        ROOT / ".impeccable/review/desktop.png": ASSETS / "admin-login-desktop.png",
        ROOT / ".impeccable/review/mobile.png": ASSETS / "admin-login-mobile.png",
    }
    for source, target in screenshots.items():
        if source.exists() and not target.exists():
            shutil.copy2(source, target)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_cell_margins(cell, top=35, start=80, bottom=35, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def set_repeat_header_footer(section):
    header = section.header
    hp = header.paragraphs[0]
    hp.text = "TRÀ SỮA NGON — BÁO CÁO TIỂU LUẬN HỆ THỐNG THƯƠNG MẠI ĐIỆN TỬ"
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.paragraph_format.space_after = Pt(0)
    for run in hp.runs:
        run.font.name = "Times New Roman"; run.font.size = Pt(9); run.font.color.rgb = RGBColor(*hex_rgb(MUTED))
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    run.font.name = "Times New Roman"; run.font.size = Pt(10)


def set_page_number_start(section, value: int):
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType"); sect_pr.append(node)
    node.set(qn("w:start"), str(value))


def add_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t"); display.text = "Mục lục sẽ được cập nhật tự động khi mở bằng Microsoft Word."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])


def add_hyperlink(paragraph, text: str, url: str):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0B513C")
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
    rpr.extend([color, underline]); run.append(rpr)
    node = OxmlElement("w:t"); node.text = text; run.append(node); hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://[^\s]+)")


def add_inline(paragraph, text: str):
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2]); run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(10.5)
        elif token.startswith("http"):
            trailing = ""
            while token and token[-1] in ".,;":
                trailing = token[-1] + trailing; token = token[:-1]
            add_hyperlink(paragraph, token, token)
            if trailing: paragraph.add_run(trailing)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def style_all_runs(paragraph, size=13, name="Times New Roman"):
    for run in paragraph.runs:
        if run.font.name != "Consolas":
            run.font.name = name
            run.font.size = Pt(size)
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def add_body_paragraph(doc, text, style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    add_inline(p, text)
    style_all_runs(p)
    return p


def add_image(doc, path: Path, alt: str):
    with Image.open(path) as im:
        px_w, px_h = im.size
    max_w, max_h = (4.20, 4.30) if path.name == "order-flow.png" else (5.00, 4.80)
    scale = min(max_w / px_w, max_h / px_h)
    width, height = px_w * scale, px_h * scale
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)
    run = p.add_run(); shape = run.add_picture(str(path), width=Inches(width), height=Inches(height))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt); doc_pr.set("title", alt)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    total_width = 6.10
    font_size = 8 if cols >= 7 else 9 if cols >= 5 else 9.2 if cols <= 2 else 10
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            cell.width = Inches(total_width / cols)
            text = row[ci] if ci < len(row) else ""
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.05; p.paragraph_format.space_after = Pt(0)
            add_inline(p, text)
            for run in p.runs:
                run.font.name = "Times New Roman"; run.font.size = Pt(font_size)
                run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
                if ri == 0: run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
            if ri == 0: set_cell_shading(cell, JADE)
            elif ri % 2 == 0: set_cell_shading(cell, "F2F6F3")
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(13)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    h1 = styles["Heading 1"]
    h1.font.name = "Times New Roman"; h1.font.size = Pt(17); h1.font.bold = True; h1.font.color.rgb = RGBColor(*hex_rgb(JADE_DARK))
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h1.paragraph_format.space_before = Pt(6); h1.paragraph_format.space_after = Pt(12); h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.page_break_before = False

    h2 = styles["Heading 2"]
    h2.font.name = "Times New Roman"; h2.font.size = Pt(14.5); h2.font.bold = True; h2.font.color.rgb = RGBColor(*hex_rgb(JADE))
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h2.paragraph_format.space_before = Pt(6); h2.paragraph_format.space_after = Pt(3); h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Times New Roman"; h3.font.size = Pt(13); h3.font.bold = True; h3.font.italic = True; h3.font.color.rgb = RGBColor(*hex_rgb(INK))
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h3.paragraph_format.space_before = Pt(8); h3.paragraph_format.space_after = Pt(4); h3.paragraph_format.keep_with_next = True

    if "Caption" not in [s.name for s in styles]:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption = styles["Caption"]
    caption.font.name = "Times New Roman"; caption.font.size = Pt(11); caption.font.italic = True; caption.font.color.rgb = RGBColor(*hex_rgb(MUTED))
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption.paragraph_format.line_spacing = 1.15; caption.paragraph_format.space_after = Pt(7); caption.paragraph_format.keep_with_next = False

    for toc_name, size, indent in (("TOC 1", 10.5, 0), ("TOC 2", 9.5, 0.35), ("TOC 3", 9, 0.7)):
        if toc_name in [s.name for s in styles]:
            toc = styles[toc_name]
            toc.font.name = "Times New Roman"; toc.font.size = Pt(size)
            toc._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            toc.paragraph_format.line_spacing = 1.0; toc.paragraph_format.space_after = Pt(0); toc.paragraph_format.left_indent = Cm(indent)

    if "Code Block" not in [s.name for s in styles]:
        styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code = styles["Code Block"]
    code.font.name = "Consolas"; code.font.size = Pt(9)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.paragraph_format.line_spacing = 1.05; code.paragraph_format.space_after = Pt(0); code.paragraph_format.left_indent = Cm(0.35); code.paragraph_format.right_indent = Cm(0.35)


def setup_sections(doc: Document):
    first = doc.sections[0]
    first.page_width = Cm(21); first.page_height = Cm(29.7)
    first.left_margin = Cm(3); first.right_margin = Cm(2); first.top_margin = Cm(2); first.bottom_margin = Cm(2)
    first.header_distance = Cm(0.8); first.footer_distance = Cm(0.9)
    return first


def add_cover(doc: Document, lines: list[str]):
    meaningful = [raw for raw in lines if raw.strip()]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    for index, raw in enumerate(meaningful):
        text = raw.strip()
        if text.startswith("# "):
            value = text[2:]
            size = 18 if "XÂY DỰNG" in value else 16
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(42 if value == "BÁO CÁO TIỂU LUẬN" else 8)
            p.paragraph_format.space_after = Pt(18 if "XÂY DỰNG" in value else 3)
            r = p.add_run(value); r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(size); r.font.color.rgb = RGBColor(*hex_rgb(JADE_DARK))
            r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        elif text.startswith("## "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text[3:]); r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(15)
            r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        else:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if index == len(meaningful) - 1:
                p.paragraph_format.space_before = Pt(48)
            add_inline(p, text.replace("  ", "")); style_all_runs(p, 13)


def build_docx():
    markdown = MARKDOWN.read_text(encoding="utf-8")
    sections = markdown.split("\n---\n")
    doc = Document()
    configure_styles(doc); setup_sections(doc)
    doc.core_properties.title = "Báo cáo tiểu luận thương mại điện tử Trà Sữa Ngon"
    doc.core_properties.subject = "Phân tích, thiết kế, xây dựng và kiểm thử website Trà Sữa Ngon"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    add_cover(doc, sections[0].splitlines())

    content_section = doc.add_section(WD_SECTION.NEW_PAGE)
    content_section.page_width = Cm(21); content_section.page_height = Cm(29.7)
    content_section.left_margin = Cm(3); content_section.right_margin = Cm(2); content_section.top_margin = Cm(2); content_section.bottom_margin = Cm(2)
    content_section.header_distance = Cm(0.8); content_section.footer_distance = Cm(0.9)
    content_section.header.is_linked_to_previous = False; content_section.footer.is_linked_to_previous = False
    set_repeat_header_footer(content_section); set_page_number_start(content_section, 1)

    for section_index, chunk in enumerate(sections[1:]):
        lines = chunk.strip("\n").splitlines()
        i = 0; in_code = False; code_lines: list[str] = []; skip_toc_body = False; first_heading_seen = False
        while i < len(lines):
            line = lines[i].rstrip()
            stripped = line.strip()
            if stripped.startswith("```"):
                if not in_code:
                    in_code = True; code_lines = []
                else:
                    p = doc.add_paragraph(style="Code Block")
                    p.paragraph_format.keep_together = True
                    for idx, code_line in enumerate(code_lines):
                        run = p.add_run(code_line)
                        if idx < len(code_lines) - 1: run.add_break()
                    p_pr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F2F6F3"); p_pr.append(shd)
                    in_code = False
                i += 1; continue
            if in_code:
                code_lines.append(line); i += 1; continue
            if not stripped:
                i += 1; continue

            if stripped.startswith("# "):
                title = stripped[2:].strip()
                p = doc.add_paragraph(title, style="Heading 1")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if not first_heading_seen and section_index > 0:
                    p.paragraph_format.page_break_before = True
                first_heading_seen = True
                skip_toc_body = title == "MỤC LỤC"
                if skip_toc_body:
                    toc_p = doc.add_paragraph(); add_toc(toc_p)
                i += 1; continue
            if skip_toc_body:
                i += 1; continue
            if stripped.startswith("## "):
                doc.add_paragraph(stripped[3:].strip(), style="Heading 2"); i += 1; continue
            if stripped.startswith("### "):
                doc.add_paragraph(stripped[4:].strip(), style="Heading 3"); i += 1; continue
            image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
            if image_match:
                path = DOCS / image_match.group(2)
                add_image(doc, path, image_match.group(1)); i += 1; continue
            if stripped.startswith("*") and stripped.endswith("*") and stripped[1:-1].startswith("Hình"):
                p = doc.add_paragraph(stripped[1:-1], style="Caption"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                i += 1; continue
            if stripped.startswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i].strip()); i += 1
                add_table(doc, parse_table(table_lines)); continue
            if re.match(r"^- ", stripped):
                p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent = Cm(0.65); p.paragraph_format.first_line_indent = Cm(-0.3)
                add_inline(p, stripped[2:]); style_all_runs(p); i += 1; continue
            if re.match(r"^\d+\. ", stripped):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.75); p.paragraph_format.first_line_indent = Cm(-0.55)
                add_inline(p, stripped); style_all_runs(p); i += 1; continue
            if stripped.startswith("> "):
                p = add_body_paragraph(doc, stripped[2:], align=WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.left_indent = Cm(0.7); p.paragraph_format.right_indent = Cm(0.7)
                p_pr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F2F6F3"); p_pr.append(shd)
                i += 1; continue
            add_body_paragraph(doc, stripped, align=WD_ALIGN_PARAGRAPH.LEFT if re.match(r"^\[\d+\]", stripped) else WD_ALIGN_PARAGRAPH.JUSTIFY)
            i += 1

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields"); settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(OUTPUT)


if __name__ == "__main__":
    make_diagrams()
    build_docx()
    print(OUTPUT)
